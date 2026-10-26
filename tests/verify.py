#!/usr/bin/env python3
"""Verify a workbook produced by docs/patch.js.

Recomputes what SPI-Generator.py would have written, then checks the generated
file cell by cell, plus the things the XML patch approach is responsible for:
macros, untouched cells, unhidden rows, and fullCalcOnLoad.

Usage: python3 tests/verify.py <actual.xlsm> <original.xlsm> <rfc.xlsm>
Prints JSON. Exit code 1 if anything is wrong.
"""
import json
import sys
import zipfile

import openpyxl

TASKS = ['NEGOTIATED ITEM PRICE', 'TOTAL QUANTITIES', 'TASK 1', 'TASK 2', 'TASK 3',
         'TASK 4', 'TASK 5', 'TASK 6', 'TASK 7', 'TASK 8', 'TASK 9', 'TASK 10']


def pay_items():
    """Same list the Python app builds in get_pay_items()."""
    import ast
    import pathlib
    tree = ast.parse(pathlib.Path("SPI-Generator.py").read_text())
    for fn in ast.walk(tree):
        if isinstance(fn, ast.FunctionDef) and fn.name == "get_pay_items":
            for node in fn.body:
                if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name):
                    return ast.literal_eval(node.value)
    raise SystemExit("could not read pay items from SPI-Generator.py")


def main():
    actual_path, original_path, rfc_path = sys.argv[1:4]
    pay = pay_items()
    problems = []

    rfc = openpyxl.load_workbook(rfc_path, read_only=True, data_only=True)["RFC"]
    expected = []
    for offset, row in enumerate(rfc.iter_rows(min_row=22, max_row=111, min_col=1, max_col=9,
                                               values_only=True)):
        if row[0] is None:
            continue
        expected.append({
            "task": "TASK " + str(row[0]),
            "pay_item": row[1],
            "units": row[7],
            "row": 11 + pay.index(row[1]),
            "col": 7 + TASKS.index("TASK " + str(row[0])),
        })

    original = openpyxl.load_workbook(original_path, keep_vba=True)
    actual = openpyxl.load_workbook(actual_path, keep_vba=True)
    sow_old = original["SOW Units"]
    sow_new = actual["SOW Units"]

    # 1. every expected change landed, with the right arithmetic
    changes = []
    for e in expected:
        old = sow_old.cell(row=e["row"], column=e["col"]).value
        old = 0 if old is None else old
        want = old + (0 if e["units"] is None else e["units"])
        got = sow_new.cell(row=e["row"], column=e["col"]).value
        changes.append({
            "address": openpyxl.utils.get_column_letter(e["col"]) + str(e["row"]),
            "task": e["task"], "pay_item": e["pay_item"],
            "old": old, "units": e["units"], "want": want, "got": got,
        })
        if got != want:
            problems.append(f"{e['task']} {e['pay_item']}: want {want}, got {got}")

    # 2. rows that got a value must be visible
    for e in expected:
        if sow_new.row_dimensions[e["row"]].hidden:
            problems.append(f"row {e['row']} is still hidden")

    # 3. untouched cells stay untouched
    collateral = 0
    for r in range(11, 503):
        for c in range(9, 19):
            if (r, c) in {(e["row"], e["col"]) for e in expected}:
                continue
            if sow_old.cell(row=r, column=c).value != sow_new.cell(row=r, column=c).value:
                collateral += 1
    if collateral:
        problems.append(f"{collateral} unrelated cells changed")

    # 4. macros survive byte for byte
    with zipfile.ZipFile(original_path) as zo, zipfile.ZipFile(actual_path) as za:
        vba_old = zo.read("xl/vbaProject.bin")
        vba_new = za.read("xl/vbaProject.bin")
        if vba_old != vba_new:
            problems.append("vbaProject.bin changed")
        missing = set(zo.namelist()) - set(za.namelist())
        if missing:
            problems.append(f"parts missing from output: {sorted(missing)}")
        wbxml = za.read("xl/workbook.xml").decode()
        if 'fullCalcOnLoad="1"' not in wbxml:
            problems.append("fullCalcOnLoad not set")

    # 5. structure intact
    if actual.sheetnames != original.sheetnames:
        problems.append("sheet list changed")

    print(json.dumps({
        "ok": not problems,
        "problems": problems,
        "changes": changes,
        "vba_bytes": len(vba_new),
        "parts": len(zipfile.ZipFile(actual_path).namelist()),
    }, indent=1))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
